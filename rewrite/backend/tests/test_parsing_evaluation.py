import hashlib
import io
import json
from pathlib import Path

import pytest
from docx import Document

from helpdesk.adapters.parsing.paragraphs import parse
from helpdesk.adapters.parsing.process import parse_file
from helpdesk.application.adjudication import adjudicate
from helpdesk.application.evaluation import Anchor, covered, percentile, write_new
from helpdesk.domain.errors import AppError


def test_paragraph_coordinates_overlap_and_docx_parsing():
    source = "甲" * 1000 + "\n\n第二段"
    chunks = parse(source.encode(), "test.md")
    assert [(c.locator["paragraph"], c.locator["start"], c.locator["end"]) for c in chunks] == [
        (1, 0, 700),
        (1, 580, 1000),
        (3, 0, 3),
    ]
    anchor = Anchor(version_id="version", paragraph=1, start=650, end=750, quote="甲" * 100)
    rows = [
        {"id": str(i), "version_id": "version", "text": c.text, "locator": c.locator}
        for i, c in enumerate(chunks)
    ]
    assert covered(anchor, rows) and not covered(anchor, rows[:1])
    document = Document()
    document.add_paragraph("DOCX 合成正文")
    output = io.BytesIO()
    document.save(output)
    assert parse(output.getvalue(), "sample.docx")[0].text == "DOCX 合成正文"


@pytest.mark.parametrize(
    "data,name",
    [(b"  \n", "empty.txt"), (b"not pdf", "bad.pdf"), (b"not zip", "bad.docx"), (b"\xff", "bad.txt")],
)
async def test_bad_document_has_no_publishable_chunks(data, name):
    with pytest.raises(AppError, match="解析失败"):
        await parse_file(data, name)


def test_human_adjudication_keeps_failures_and_pending(tmp_path):
    report = tmp_path / "report"
    report.mkdir()
    cases = {"expected_behavior": "answer"}
    write_new(
        report / "blind-review.json",
        [
            {"review_id": "R1", "case": cases, "status": "error"},
            {"review_id": "R2", "case": cases, "status": "completed"},
        ],
    )
    write_new(report / "private-mapping.json", {"R1": "one", "R2": "two"})
    write_new(
        report / "results.json",
        [
            {
                "key": "one",
                "case_id": "case",
                "strategy": "baseline",
                "variant": "D",
                "status": "error",
                "calls": 1,
            },
            {
                "key": "two",
                "case_id": "case",
                "strategy": "split-v1",
                "variant": "D",
                "status": "completed",
                "calls": 4,
            },
        ],
    )
    decision = tmp_path / "decisions.jsonl"
    rows = [
        {
            "review_id": "R1",
            "correct": True,
            "unsupported": False,
            "omitted_requirements": [],
            "reviewer": "synthetic",
            "reason": "test",
        }
    ]
    decision.write_text(json.dumps(rows[0]), encoding="utf-8")
    with pytest.raises(ValueError, match="中断"):
        adjudicate(report, decision, tmp_path / "invalid")
    rows[0]["correct"] = False
    decision.write_text(json.dumps(rows[0]), encoding="utf-8")
    adjudicate(report, decision, tmp_path / "partial")
    result = json.loads((tmp_path / "partial/summary.json").read_text(encoding="utf-8"))
    assert result["pending_recorded"] == 1 and result["pairs"] == []
    assert percentile([], 0.95) is None and percentile([1, 2, 3, 4], 0.95) == 4


def test_prompt_assets_match_frozen_manifest():
    from helpdesk.graph.baseline import PROMPTS

    manifest = json.loads(
        (Path(__file__).resolve().parents[2] / "contracts/prompt-manifest.json").read_text(encoding="utf-8")
    )
    assert manifest["sha256"] == {
        key: hashlib.sha256(text.encode()).hexdigest() for key, text in PROMPTS.items()
    }
